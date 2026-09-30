import calendar
import hashlib
from io import BytesIO
import json
import os
import random
import re
import time as sleep_time
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from threading import Lock
from typing import Iterable

from dateutil import parser as date_parser
from googleapiclient.errors import HttpError
from openpyxl.cell.rich_text import CellRichText, TextBlock
from googleapiclient.http import MediaIoBaseDownload
from openpyxl import load_workbook


COURSE_RE = re.compile(r"(?P<code>26[A-Z0-9]+)\s*:\s*(?P<subject>.+?)\s*\[(?P<section>[A-Z])\]")
COURSE_START_RE = re.compile(r"(?=26[A-Z0-9]+\s*:)")
TIME_RE = re.compile(
    r"(?P<start>\d{1,2}[.:]\d{2})\s*(?P<start_ampm>[AP]M)?\s*[-–]\s*"
    r"(?P<end>\d{1,2}[.:]\d{2})\s*(?P<end_ampm>[AP]M)?",
    re.IGNORECASE,
)
WEEK_RE = re.compile(
    r"\((?P<start_day>\d{1,2})\s*(?P<start_month>[A-Za-z]+)?\s*[-–]\s*"
    r"(?P<end_day>\d{1,2})\s*(?P<end_month>[A-Za-z]+)\)"
)
DAY_INDEX = {
    "MON": 0,
    "MONDAY": 0,
    "TUE": 1,
    "TUESDAY": 1,
    "WED": 2,
    "WEDNESDAY": 2,
    "THU": 3,
    "THURSDAY": 3,
    "FRI": 4,
    "FRIDAY": 4,
    "SAT": 5,
    "SATURDAY": 5,
    "SUN": 6,
    "SUNDAY": 6,
}
RETRYABLE_STATUSES = {429, 500, 502, 503, 504}
RETRYABLE_REASONS = {"rateLimitExceeded", "userRateLimitExceeded", "quotaExceeded", "backendError"}
CONFLICT_STATUS = 409
TIMETABLE_CACHE_TTL_SECONDS = int(os.environ.get("TIMETABLE_CACHE_TTL_SECONDS", "300"))
_timetable_cache_lock = Lock()
_timetable_cache: dict[str, dict] = {}


@dataclass(frozen=True)
class CourseEnrollment:
    course_code: str
    subject: str
    section: str


@dataclass(frozen=True)
class Student:
    student_id: str
    name: str
    courses: list[CourseEnrollment] = field(default_factory=list)


@dataclass(frozen=True)
class TimetableEvent:
    date: str
    start_time: str
    end_time: str
    course_code: str
    subject: str
    section: str
    location: str = ""
    professor: str = ""
    source_sheet: str = ""
    cancelled: bool = False

    @property
    def start_datetime(self) -> str:
        return f"{self.date}T{self.start_time}:00"

    @property
    def end_datetime(self) -> str:
        return f"{self.date}T{self.end_time}:00"


class SectionRepository:
    def __init__(self, workbook_path: str):
        self.workbook_path = workbook_path
        self._students = self._load_students()

    def _load_students(self) -> dict[str, Student]:
        wb = load_workbook(self.workbook_path, read_only=True, data_only=True)
        ws = wb.active
        headers = [str(cell.value).strip() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        indexes = {name: headers.index(name) for name in ["Section", "CourseCode", "StudentId", "Name"]}
        grouped: dict[str, dict] = {}

        for row in ws.iter_rows(min_row=2, values_only=True):
            section_label = str(row[indexes["Section"]] or "").strip()
            course_code = str(row[indexes["CourseCode"]] or "").strip()
            student_id = str(row[indexes["StudentId"]] or "").strip()
            name = str(row[indexes["Name"]] or "").strip()
            if not section_label or not course_code or not student_id or not name:
                continue
            parsed = parse_section_label(section_label)
            if not parsed:
                continue
            course = CourseEnrollment(course_code=course_code, subject=parsed["subject"], section=parsed["section"])
            grouped.setdefault(student_id.upper(), {"student_id": student_id, "name": name, "courses": []})
            grouped[student_id.upper()]["courses"].append(course)

        return {
            key: Student(
                student_id=value["student_id"],
                name=value["name"],
                courses=dedupe_courses(value["courses"]),
            )
            for key, value in grouped.items()
        }

    def find_student(self, identifier: str) -> Student:
        clean = normalize(identifier)
        if not clean:
            raise ValueError("Enter a student name or student ID.")
        if clean.upper() in self._students:
            return self._students[clean.upper()]

        exact = [student for student in self._students.values() if normalize(student.name) == clean]
        if len(exact) == 1:
            return exact[0]

        partial = [student for student in self._students.values() if clean in normalize(student.name)]
        if len(partial) == 1:
            return partial[0]
        if len(partial) > 1:
            names = ", ".join(f"{s.name} ({s.student_id})" for s in partial[:8])
            raise ValueError(f"Multiple students matched. Use student ID. Matches: {names}")

        raise ValueError("No matching student found.")


class TimetableService:
    def __init__(self, spreadsheet_id: str, credentials, timezone: str = "Asia/Kolkata"):
        from googleapiclient.discovery import build

        self.spreadsheet_id = spreadsheet_id
        self.credentials = credentials
        self.timezone = timezone
        self.service = build("sheets", "v4", credentials=credentials)

    def match_events(self, enrollments: Iterable[CourseEnrollment]) -> list[TimetableEvent]:
        wanted = {(course.course_code.upper(), course.section.upper()) for course in enrollments}
        all_events = self.read_events()
        return [
            event
            for event in all_events
            if (event.course_code.upper(), event.section.upper()) in wanted
            and datetime.fromisoformat(event.end_datetime) >= datetime.now()
        ]

    def read_events(self) -> list[TimetableEvent]:
        cached = get_timetable_cache(self.spreadsheet_id)
        if cached is not None:
            return cached

        if TIMETABLE_CACHE_TTL_SECONDS > 0:
            with _timetable_cache_lock:
                cached = get_timetable_cache_unlocked(self.spreadsheet_id)
                if cached is not None:
                    return cached
                return set_timetable_cache_unlocked(self.spreadsheet_id, self._read_events_uncached())

        return self._read_events_uncached()

    def _read_events_uncached(self) -> list[TimetableEvent]:
        try:
            meta = execute_google_request(
                self.service.spreadsheets().get(spreadsheetId=self.spreadsheet_id),
                "read spreadsheet metadata",
            )
            sheet_titles = [sheet["properties"]["title"] for sheet in meta.get("sheets", [])]
            selected = self._select_sheet_titles(sheet_titles)
            events: list[TimetableEvent] = []
            for title in selected:
                values = execute_google_request(
                    self.service.spreadsheets()
                    .values()
                    .get(spreadsheetId=self.spreadsheet_id, range=f"'{title}'!A1:H40"),
                    f"read timetable tab {title}",
                ).get("values", [])
                events.extend(self._parse_grid(title, values))
            return sorted(events, key=lambda event: (event.date, event.start_time, event.course_code))
        except Exception as exc:
            if not is_office_file_error(exc):
                raise
            return self._read_events_from_drive_xlsx()

    def _read_events_from_drive_xlsx(self) -> list[TimetableEvent]:
        from googleapiclient.discovery import build

        drive = build("drive", "v3", credentials=self.credentials)
        request = drive.files().get_media(fileId=self.spreadsheet_id)
        buffer = BytesIO()
        downloader = MediaIoBaseDownload(buffer, request)
        done = False
        while not done:
            _, done = execute_google_download_chunk(downloader, "download timetable workbook")
        buffer.seek(0)
        if os.environ.get("DEBUG_TIMETABLE_XLSX"):
            with open(os.environ["DEBUG_TIMETABLE_XLSX"], "wb") as debug_file:
                debug_file.write(buffer.getvalue())

        workbook = load_workbook(buffer, read_only=False, data_only=True, rich_text=True)
        selected = self._select_sheet_titles(workbook.sheetnames)
        events: list[TimetableEvent] = []
        for title in selected:
            worksheet = workbook[title]
            values = []
            for row in worksheet.iter_rows(min_row=1, max_row=40, max_col=8):
                values.append([cell_to_grid_value(cell) for cell in row])
            events.extend(self._parse_grid(title, values))
        return sorted(events, key=lambda event: (event.date, event.start_time, event.course_code))

    def _select_sheet_titles(self, titles: list[str]) -> list[str]:
        dated = []
        today = date.today()
        for title in titles:
            week_start = parse_week_start(title, today.year)
            if week_start:
                dated.append((week_start, title))
        if not dated:
            raise ValueError("Could not find week dates in any timetable tab name.")

        dated.sort()
        current_or_future = [(start, title) for start, title in dated if start + timedelta(days=6) >= today]
        return [title for _, title in current_or_future] if current_or_future else [title for _, title in dated]

    def _parse_grid(self, sheet_title: str, values: list[list[str]]) -> list[TimetableEvent]:
        if not values:
            return []
        week_start = parse_week_start(sheet_title, date.today().year)
        if not week_start:
            week_start = infer_week_start_from_headers(values)
        if not week_start:
            raise ValueError(f"Could not infer week dates from sheet '{sheet_title}'.")

        header = values[0]
        time_columns = {
            col_idx: parsed
            for col_idx, cell in enumerate(header)
            if col_idx > 0 and (parsed := parse_time_range(grid_text(cell)))
        }
        events: list[TimetableEvent] = []
        current_day_index = None
        for row in values[1:]:
            if not row:
                continue
            day_label = grid_text(row[0] if len(row) > 0 else "").strip().upper()
            if day_label in DAY_INDEX:
                current_day_index = DAY_INDEX[day_label]
            if current_day_index is None:
                continue

            for col_idx, start_end in time_columns.items():
                if col_idx >= len(row):
                    continue
                cell_value = row[col_idx]
                cell_text = grid_text(cell_value).strip()
                if not cell_text:
                    continue
                for parsed in parse_timetable_cell(cell_text, grid_cancelled_text(cell_value), grid_cancelled(cell_value)):
                    event_date = week_start + timedelta(days=current_day_index)
                    events.append(
                        TimetableEvent(
                            date=event_date.isoformat(),
                            start_time=start_end[0],
                            end_time=start_end[1],
                            course_code=parsed["course_code"],
                            subject=parsed["subject"],
                            section=parsed["section"],
                            professor=parsed.get("professor", ""),
                            location=parsed.get("location", ""),
                            source_sheet=sheet_title,
                            cancelled=parsed.get("cancelled", False),
                        )
                    )
        return events


class CalendarSyncService:
    def __init__(self, credentials, timezone: str = "Asia/Kolkata", color_id: str = "10"):
        from googleapiclient.discovery import build

        self.service = build("calendar", "v3", credentials=credentials)
        self.timezone = timezone
        self.color_id = color_id

    def upsert_events(self, student: Student, events: list[TimetableEvent]) -> dict[str, int]:
        created = 0
        updated = 0
        for event in events:
            agent_id = event_agent_id(student, event)
            body = self._event_body(student, event, agent_id)
            event_id = calendar_event_id(agent_id)

            try:
                execute_google_request(
                    self.service.events().update(calendarId="primary", eventId=event_id, body=body),
                    f"update calendar event {agent_id}",
                    passthrough_statuses={404},
                )
                updated += 1
                continue
            except HttpError as exc:
                if not is_not_found(exc):
                    raise

            body["id"] = event_id
            try:
                execute_google_request(
                    self.service.events().insert(calendarId="primary", body=body),
                    f"insert calendar event {agent_id}",
                )
                created += 1
            except HttpError as exc:
                if not is_conflict(exc):
                    raise
                execute_google_request(
                    self.service.events().update(calendarId="primary", eventId=event_id, body=body),
                    f"update existing calendar event {agent_id}",
                )
                updated += 1
        return {"created": created, "updated": updated}

    def delete_events(self, student: Student, events: list[TimetableEvent]) -> dict[str, int]:
        deleted = 0
        missing = 0
        for event in events:
            agent_id = event_agent_id(student, event)
            event_id = calendar_event_id(agent_id)
            try:
                execute_google_request(
                    self.service.events().delete(calendarId="primary", eventId=event_id),
                    f"delete calendar event {agent_id}",
                    passthrough_statuses={404, 410},
                )
                deleted += 1
            except HttpError as exc:
                if is_not_found(exc):
                    missing += 1
                    continue
                raise
        return {"deleted": deleted, "missing": missing}

    def _event_body(self, student: Student, event: TimetableEvent, agent_id: str) -> dict:
        summary = f"{event.course_code}: {event.subject} [{event.section}]"
        reminders = {"useDefault": False, "overrides": [{"method": "popup", "minutes": 10}]}
        if event.cancelled:
            summary = f"Cancelled: {summary}"
            reminders = {"useDefault": False, "overrides": []}
        return {
            "summary": summary,
            "location": event.location,
            "description": (
                f"For {student.name} ({student.student_id}).\\n"
                f"Status: {'Cancelled in live timetable' if event.cancelled else 'Scheduled'}\\n"
                f"Professor: {event.professor or 'Not listed'}\\n"
                f"Source: live timetable sheet, tab {event.source_sheet}."
            ),
            "start": {"dateTime": event.start_datetime, "timeZone": self.timezone},
            "end": {"dateTime": event.end_datetime, "timeZone": self.timezone},
            "colorId": self.color_id,
            "reminders": reminders,
            "extendedProperties": {"private": {"iimAgentId": agent_id}},
        }


def normalize(value: str) -> str:
    return re.sub(r"\\s+", " ", str(value or "").strip()).lower()


def dedupe_courses(courses: list[CourseEnrollment]) -> list[CourseEnrollment]:
    seen = set()
    result = []
    for course in courses:
        key = (course.course_code.upper(), course.section.upper())
        if key not in seen:
            seen.add(key)
            result.append(course)
    return result


def parse_section_label(value: str) -> dict | None:
    match = COURSE_RE.search(value)
    if not match:
        return None
    return {
        "course_code": match.group("code").strip(),
        "subject": match.group("subject").strip(),
        "section": match.group("section").strip(),
    }


def parse_time_range(value: str) -> tuple[str, str] | None:
    match = TIME_RE.search(str(value or ""))
    if not match:
        return None
    start_ampm = match.group("start_ampm") or match.group("end_ampm")
    end_ampm = match.group("end_ampm") or start_ampm
    return normalize_time(match.group("start"), start_ampm), normalize_time(match.group("end"), end_ampm)


def normalize_time(value: str, ampm: str | None = None) -> str:
    clean = value.strip().replace(".", ":")
    hour_text, minute_text = clean.split(":", 1)
    hour = int(hour_text)
    minute = int(minute_text[:2])
    if ampm:
        marker = ampm.upper()
        if marker == "PM" and hour != 12:
            hour += 12
        elif marker == "AM" and hour == 12:
            hour = 0
    return f"{hour:02d}:{minute:02d}"


def parse_timetable_cell(cell: str, cancelled_text: str = "", cell_cancelled: bool = False) -> list[dict]:
    text = str(cell or "").strip()
    if not text:
        return []

    results = []
    for block in [part.strip() for part in COURSE_START_RE.split(text) if part.strip()]:
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        joined = " ".join(lines)
        match = COURSE_RE.search(joined)
        if not match:
            continue

        professor = ""
        location = ""
        for line in lines[1:]:
            lowered = line.lower()
            if lowered.startswith("prof"):
                professor = line
            elif line.startswith("(") and line.endswith(")"):
                location = line.strip("()")
            elif not location:
                location = line

        results.append(
            {
                "course_code": match.group("code").strip(),
                "subject": match.group("subject").strip(),
                "section": match.group("section").strip(),
                "professor": professor,
                "location": location,
                "cancelled": cell_cancelled or block_is_cancelled(block, match, cancelled_text),
            }
        )
    return results


def cell_to_grid_value(cell) -> dict:
    is_rich_text = isinstance(cell.value, CellRichText)
    text, cancelled_text = rich_text_parts(cell.value, cell.font)
    if not text:
        text = "" if cell.value is None else str(cell.value)
    cell_cancelled = not is_rich_text and is_struck(cell.font)
    return {
        "text": text,
        "cancelled": cell_cancelled,
        "cancelled_text": cancelled_text,
    }


def rich_text_parts(value, default_font=None) -> tuple[str, str]:
    if not isinstance(value, CellRichText):
        return "", ""

    text_parts = []
    cancelled_parts = []
    for part in value:
        if isinstance(part, TextBlock):
            text_parts.append(part.text)
            if is_struck(part.font):
                cancelled_parts.append(part.text)
        else:
            part_text = str(part)
            text_parts.append(part_text)
            if is_struck(default_font):
                cancelled_parts.append(part_text)
    return "".join(text_parts), "".join(cancelled_parts)


def is_struck(font) -> bool:
    return bool(getattr(font, "strike", False))


def grid_text(value) -> str:
    if isinstance(value, dict):
        return str(value.get("text") or "")
    return str(value or "")


def grid_cancelled(value) -> bool:
    return bool(isinstance(value, dict) and value.get("cancelled"))


def grid_cancelled_text(value) -> str:
    if isinstance(value, dict):
        return str(value.get("cancelled_text") or "")
    return ""


def block_is_cancelled(block: str, match, cancelled_text: str) -> bool:
    if not cancelled_text:
        return False
    event_key = (match.group("code").strip().upper(), match.group("section").strip().upper())
    for cancelled_block in [part.strip() for part in COURSE_START_RE.split(cancelled_text) if part.strip()]:
        cancelled_match = COURSE_RE.search(" ".join(line.strip() for line in cancelled_block.splitlines()))
        if not cancelled_match:
            continue
        cancelled_key = (
            cancelled_match.group("code").strip().upper(),
            cancelled_match.group("section").strip().upper(),
        )
        if cancelled_key == event_key:
            return True
    return False


def parse_week_start(title: str, year: int) -> date | None:
    match = WEEK_RE.search(title)
    if not match:
        return None
    start_month = match.group("start_month") or match.group("end_month")
    start_text = f"{match.group('start_day')} {start_month}"
    start_text = re.sub(r"\\s+", " ", start_text.strip())
    parsed = date_parser.parse(f"{start_text} {year}", dayfirst=True, fuzzy=True).date()
    return parsed


def infer_week_start_from_headers(values: list[list[str]]) -> date | None:
    return None


def get_timetable_cache(spreadsheet_id: str) -> list[TimetableEvent] | None:
    if TIMETABLE_CACHE_TTL_SECONDS <= 0:
        return None
    with _timetable_cache_lock:
        return get_timetable_cache_unlocked(spreadsheet_id)


def get_timetable_cache_unlocked(spreadsheet_id: str) -> list[TimetableEvent] | None:
    cached = _timetable_cache.get(spreadsheet_id)
    if not cached:
        return None
    if sleep_time.time() - cached["loaded_at"] > TIMETABLE_CACHE_TTL_SECONDS:
        _timetable_cache.pop(spreadsheet_id, None)
        return None
    return list(cached["events"])


def set_timetable_cache_unlocked(spreadsheet_id: str, events: list[TimetableEvent]) -> list[TimetableEvent]:
    _timetable_cache[spreadsheet_id] = {"loaded_at": sleep_time.time(), "events": list(events)}
    return events


def execute_google_request(request, operation: str, max_attempts: int = 5, passthrough_statuses: set[int] | None = None):
    passthrough_statuses = passthrough_statuses or set()
    for attempt in range(max_attempts):
        try:
            return request.execute()
        except HttpError as exc:
            if is_conflict(exc) or getattr(exc.resp, "status", None) in passthrough_statuses:
                raise
            if not should_retry_google_error(exc) or attempt == max_attempts - 1:
                raise RuntimeError(google_error_message(exc, operation)) from exc
            sleep_time.sleep(retry_delay(attempt, exc))
    raise RuntimeError(f"Google API request failed while trying to {operation}.")


def execute_google_download_chunk(downloader, operation: str, max_attempts: int = 5):
    for attempt in range(max_attempts):
        try:
            return downloader.next_chunk()
        except HttpError as exc:
            if not should_retry_google_error(exc) or attempt == max_attempts - 1:
                raise RuntimeError(google_error_message(exc, operation)) from exc
            sleep_time.sleep(retry_delay(attempt, exc))
    raise RuntimeError(f"Google API request failed while trying to {operation}.")


def should_retry_google_error(exc: HttpError) -> bool:
    status = getattr(exc.resp, "status", None)
    if status in RETRYABLE_STATUSES:
        return True
    return status == 403 and google_error_reason(exc) in RETRYABLE_REASONS


def retry_delay(attempt: int, exc: HttpError) -> float:
    retry_after = getattr(exc.resp, "get", lambda _key: None)("retry-after")
    if retry_after:
        try:
            return min(float(retry_after), 30.0)
        except ValueError:
            pass
    return min((2**attempt) + random.uniform(0, 0.5), 30.0)


def google_error_reason(exc: HttpError) -> str:
    try:
        payload = json.loads(exc.content.decode("utf-8"))
    except Exception:
        return ""
    errors = payload.get("error", {}).get("errors", [])
    if errors:
        return str(errors[0].get("reason", ""))
    return str(payload.get("error", {}).get("status", ""))


def google_error_message(exc: HttpError, operation: str) -> str:
    status = getattr(exc.resp, "status", "unknown")
    reason = google_error_reason(exc)
    detail = f" {reason}" if reason else ""
    return f"Google API error while trying to {operation}: HTTP {status}{detail}."


def is_conflict(exc: HttpError) -> bool:
    return getattr(exc.resp, "status", None) == CONFLICT_STATUS


def is_not_found(exc: HttpError) -> bool:
    return getattr(exc.resp, "status", None) in {404, 410}


def is_office_file_error(exc: Exception) -> bool:
    current = exc
    while current:
        if "Office file" in str(current):
            return True
        current = getattr(current, "__cause__", None)
    return False


def event_agent_id(student: Student, event: TimetableEvent) -> str:
    raw = "|".join(
        [
            student.student_id.upper(),
            event.date,
            event.start_time,
            event.end_time,
            event.course_code.upper(),
            event.section.upper(),
        ]
    )
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:32]


def calendar_event_id(agent_id: str) -> str:
    return f"iim{agent_id}"
