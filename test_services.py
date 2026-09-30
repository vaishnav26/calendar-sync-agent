import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font

from services import cell_to_grid_value, grid_cancelled, parse_timetable_cell


COURSE_TEXT = "26MBA401: Digital Transformation [A]\nProf. Example\n(CR-1)"


class CancellationFormattingTests(unittest.TestCase):
    def load_cell(self, font: Font, value=COURSE_TEXT):
        workbook = Workbook()
        cell = workbook.active["B2"]
        cell.value = value
        cell.font = font

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "timetable.xlsx"
            workbook.save(path)
            loaded = load_workbook(path, read_only=False, data_only=True, rich_text=True)
            return loaded.active["B2"]

    def test_whole_cell_strikethrough_marks_class_cancelled(self):
        value = cell_to_grid_value(self.load_cell(Font(strike=True)))
        parsed = parse_timetable_cell(value["text"], value["cancelled_text"], grid_cancelled(value))

        self.assertTrue(value["cancelled"])
        self.assertTrue(parsed[0]["cancelled"])

    def test_red_text_without_strikethrough_stays_scheduled(self):
        value = cell_to_grid_value(self.load_cell(Font(color="FFFF0000", strike=False)))
        parsed = parse_timetable_cell(value["text"], value["cancelled_text"], grid_cancelled(value))

        self.assertFalse(value["cancelled"])
        self.assertFalse(parsed[0]["cancelled"])

    def test_visible_rich_text_overrides_stale_cell_strikethrough(self):
        rich_value = CellRichText([TextBlock(InlineFont(strike=False), COURSE_TEXT)])
        value = cell_to_grid_value(self.load_cell(Font(strike=True), rich_value))
        parsed = parse_timetable_cell(value["text"], value["cancelled_text"], grid_cancelled(value))

        self.assertFalse(value["cancelled"])
        self.assertFalse(parsed[0]["cancelled"])

    def test_visible_rich_text_strikethrough_marks_class_cancelled(self):
        rich_value = CellRichText([TextBlock(InlineFont(strike=True), COURSE_TEXT)])
        value = cell_to_grid_value(self.load_cell(Font(strike=False), rich_value))
        parsed = parse_timetable_cell(value["text"], value["cancelled_text"], grid_cancelled(value))

        self.assertFalse(value["cancelled"])
        self.assertTrue(parsed[0]["cancelled"])


if __name__ == "__main__":
    unittest.main()
